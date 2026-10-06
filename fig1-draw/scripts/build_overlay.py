#!/usr/bin/env python3
"""Build an editable PPTX from an AI PNG plus a region manifest.

Text becomes native text boxes. Simple geometry is traced to SVG and kept
only when a re-raster matches the crop. Depictions stay on the base image.
"""

from __future__ import annotations

import argparse
import io
import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

INSTALL = "python3 -m venv .venv && .venv/bin/pip install pillow python-pptx vtracer pymupdf"
PX = 9525  # 1 px = 1/96 in = 0.75 pt
# ponytail: fixed tolerances, STROKE_TOL also bounds the traced color; retune if a faithful real trace is rejected
STROKE_TOL = 36
MIN_COVERAGE = 0.90
MAX_SPILL = 0.20


def require_libs():
    missing = []
    for name in ("PIL", "pptx", "vtracer", "fitz"):
        try:
            __import__(name)
        except ImportError:
            missing.append(name)
    if missing:
        sys.exit(f"缺少 {', '.join(missing)}。空环境先执行:\n{INSTALL}")


def installed_fonts():
    try:
        raw = subprocess.check_output(["fc-list", ":", "family"], text=True, errors="replace")
    except FileNotFoundError:
        sys.exit("需要 fontconfig 的 fc-list 才能核对字体。")
    fams = set()
    for line in raw.splitlines():
        for part in line.split(","):
            part = part.strip()
            if part:
                fams.add(part)
    return sorted(fams)


def die(message, fonts=None):
    if fonts is not None:
        message += "\n已安装字体:\n" + "\n".join(fonts)
    sys.exit(message)


def load_manifest(path: Path):
    data = json.loads(path.read_text())
    if not isinstance(data, dict):
        die("清单必须是 JSON 对象")
    src = data.get("source_image")
    if not isinstance(src, str) or not src:
        die("source_image 必填")
    image = (path.parent / src).resolve()
    if not image.is_file():
        die(f"找不到源图 {image}")
    return data, image


def rect_of(value, size, where):
    if not isinstance(value, list) or len(value) != 4 or not all(isinstance(n, (int, float)) for n in value):
        die(f"{where} 的 rect 需要四个数")
    x, y, w, h = value
    width, height = size
    if x < 0 or y < 0 or w <= 0 or h <= 0 or x + w > width or y + h > height:
        die(f"{where} 的 rect 超出画面 {size}")
    # Caller sizes the rect. Glyphs are not measured here.
    return [int(x), int(y), int(w), int(h)]


def hex_color(value, where):
    if not isinstance(value, str) or len(value) != 7 or not value.startswith("#"):
        die(f"{where} 的颜色需要 #RRGGBB")
    try:
        int(value[1:], 16)
    except ValueError:
        die(f"{where} 的颜色需要 #RRGGBB")
    return value


def validate(data, size, fonts):
    ids = set()
    texts = data.get("texts") or []
    vectors = data.get("vectors") or []
    masks = data.get("masks") or []
    if not texts and not vectors:
        die("texts 和 vectors 至少一个非空")
    known = {name.lower() for name in fonts}

    def take_id(obj, where):
        item_id = obj.get("id")
        if not isinstance(item_id, str) or not item_id.strip() or item_id in ids:
            die(f"{where} 需要唯一的 id")
        ids.add(item_id)
        return item_id

    for mask in masks:
        take_id(mask, "mask")
        rect_of(mask.get("rect"), size, mask["id"])
        hex_color(mask.get("fill"), mask["id"])
    for text in texts:
        take_id(text, "text")
        rect_of(text.get("rect"), size, text["id"])
        runs = text.get("runs")
        if runs is not None:
            if not isinstance(runs, list) or not runs:
                die(f"{text['id']} 的 runs 不能为空")
            chunks = []
            for run in runs:
                if not isinstance(run, dict) or not isinstance(run.get("text"), str):
                    die(f"{text['id']} 的 run 需要 text")
                font = run.get("font") or text.get("font")
                if not isinstance(font, str) or not font.strip():
                    die(f"{text['id']} 缺少字体", fonts)
                if font.lower() not in known:
                    die(f"未安装字体 {font}", fonts)
                chunks.append(run["text"])
            if text.get("text") not in (None, "".join(chunks)):
                die(f"{text['id']} 的 text 与 runs 不一致")
        else:
            if not isinstance(text.get("text"), str) or not text["text"]:
                die(f"{text['id']} 需要 text 或 runs")
            font = text.get("font")
            if not isinstance(font, str) or not font.strip():
                die(f"{text['id']} 缺少字体", fonts)
            if font.lower() not in known:
                die(f"未安装字体 {font}", fonts)
        if not isinstance(text.get("font_size"), (int, float)) or text["font_size"] <= 0:
            die(f"{text['id']} 需要正的 font_size（源像素）")
        hex_color(text.get("color", "#142447"), text["id"])
        if text.get("align", "left") not in ("left", "center", "right", "justify"):
            die(f"{text['id']} 的 align 无效")
        if text.get("vertical_align", "middle") not in ("top", "middle", "bottom"):
            die(f"{text['id']} 的 vertical_align 无效")
    for vec in vectors:
        take_id(vec, "vector")
        rect_of(vec.get("rect"), size, vec["id"])
        kind = vec.get("kind", "geometry")
        if not isinstance(kind, str) or not kind.strip():
            die(f"{vec['id']} 的 kind 需要是字符串")
    return masks, texts, vectors


def dist(a, b):
    return abs(a[0] - b[0]) + abs(a[1] - b[1]) + abs(a[2] - b[2])


def border_bg(im):
    rgb = im.convert("RGB")
    px = rgb.load()
    w, h = rgb.size
    cols = [px[x, 0] for x in range(w)]
    if h > 1:
        cols.extend(px[x, h - 1] for x in range(w))
    for y in range(1, max(h - 1, 1)):
        cols.append(px[0, y])
        if w > 1:
            cols.append(px[w - 1, y])
    cols.sort()
    return cols[len(cols) // 2]


def knockout(im, bg):
    out = im.convert("RGBA")
    px = out.load()
    w, h = out.size
    for y in range(h):
        for x in range(w):
            r, g, b, _a = px[x, y]
            if dist((r, g, b), bg) <= STROKE_TOL:
                px[x, y] = (0, 0, 0, 0)
            else:
                px[x, y] = (r, g, b, 255)
    return out


def stroke_mask(im, bg):
    rgb = im.convert("RGB")
    px = rgb.load()
    w, h = rgb.size
    return [[dist(px[x, y], bg) > STROKE_TOL for x in range(w)] for y in range(h)]


def trace_svg(crop):
    import vtracer

    bg = border_bg(crop)
    plate = knockout(crop, bg)
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "crop.png"
        dst = Path(tmp) / "crop.svg"
        plate.save(src)
        vtracer.convert_image_to_svg_py(
            str(src),
            str(dst),
            colormode="color",
            hierarchical="stacked",
            mode="spline",
            filter_speckle=1,
            color_precision=6,
            layer_difference=16,
        )
        svg = dst.read_text()
    if "<path" not in svg:
        return bg, None, "描摹结果没有路径"
    return bg, svg, None


def rasterize(svg, size):
    import fitz
    from PIL import Image

    doc = fitz.open(stream=svg.encode(), filetype="svg")
    page = doc[0]
    w, h = size
    mat = fitz.Matrix(w / page.rect.width, h / page.rect.height)
    pix = page.get_pixmap(matrix=mat, alpha=True)
    im = Image.frombytes("RGBA", (pix.width, pix.height), pix.samples)
    if im.size != size:
        im = im.resize(size, Image.Resampling.NEAREST)
    return im


def composite(im, bg):
    from PIL import Image

    base = Image.new("RGB", im.size, bg)
    base.paste(im.convert("RGBA"), mask=im.convert("RGBA").getchannel("A"))
    return base


def matches(original, rendered, bg):
    left = stroke_mask(original, bg)
    right = stroke_mask(rendered, bg)
    src = original.convert("RGB").load()
    out = rendered.convert("RGB").load()
    base = extra = 0
    drift = []
    for y, (row_a, row_b) in enumerate(zip(left, right)):
        for x, (a, b) in enumerate(zip(row_a, row_b)):
            if a:
                base += 1
                if b:
                    drift.append(dist(src[x, y], out[x, y]))
            elif b:
                extra += 1
    if base == 0:
        return False, "区域内没有可描的笔画"
    coverage = len(drift) / base
    spill = extra / base
    color = sorted(drift)[len(drift) // 2] if drift else 0
    ok = coverage >= MIN_COVERAGE and spill <= MAX_SPILL and color <= STROKE_TOL
    return ok, f"coverage {coverage:.2f} spill {spill:.2f} color {color}"


def emu(px):
    from pptx.util import Emu

    return Emu(int(round(px * PX)))


def add_mask(slide, item_id, rect, fill):
    from lxml import etree
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

    x, y, w, h = rect
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, emu(x), emu(y), emu(w), emu(h))
    shape.name = f"MASK {item_id}"
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill[1:])
    shape.line.fill.background()
    # The default theme style puts a drop shadow on every rectangle.
    sp = shape._element
    style = sp.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if style is not None:
        sp.remove(style)
    sp_pr = sp.find("{http://schemas.openxmlformats.org/presentationml/2006/main}spPr")
    etree.SubElement(sp_pr, "{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst")


def add_text(slide, text):
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Pt

    x, y, w, h = text["rect"]
    box = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    box.name = f"TEXT {text['id']}"
    frame = box.text_frame
    frame.word_wrap = False
    frame.auto_size = None
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[
        text.get("vertical_align", "middle")
    ]
    align = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}[
        text.get("align", "left")
    ]
    runs = text["runs"] if text.get("runs") else [{"text": text["text"]}]
    lines = [[]]
    for run in runs:
        parts = run["text"].split("\n")
        for i, part in enumerate(parts):
            if i:
                lines.append([])
            lines[-1].append({**run, "text": part})
    for i, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.clear()
        for run in line:
            chunk = paragraph.add_run()
            chunk.text = run["text"]
            font = chunk.font
            font.name = run.get("font") or text["font"]
            font.size = Pt(float(run.get("font_size") or text["font_size"]) * 0.75)
            font.bold = bool(run.get("bold", text.get("bold", False)))
            font.italic = bool(run.get("italic", text.get("italic", False)))
            font.color.rgb = RGBColor.from_string((run.get("color") or text.get("color", "#142447"))[1:])


def add_svg_picture(slide, png_bytes, svg_bytes, rect, name):
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.opc.package import Part

    x, y, w, h = rect
    picture = slide.shapes.add_picture(io.BytesIO(png_bytes), emu(x), emu(y), emu(w), emu(h))
    picture.name = name
    package = slide.part.package
    partname = package.next_partname("/ppt/media/image%d.svg")
    svg_part = Part(partname, "image/svg+xml", package, svg_bytes)
    rid = slide.part.relate_to(svg_part, RT.IMAGE)
    blip = picture._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
    ext_lst = etree.SubElement(blip, "{http://schemas.openxmlformats.org/drawingml/2006/main}extLst")
    ext = etree.SubElement(ext_lst, "{http://schemas.openxmlformats.org/drawingml/2006/main}ext")
    ext.set("uri", "{96DAC541-7B7A-43D3-8B79-37D633B846F1}")
    svg_blip = etree.SubElement(ext, "{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip")
    svg_blip.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed", rid)


def build(manifest_path: Path, out_path: Path):
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Emu

    data, image_path = load_manifest(manifest_path)
    image = Image.open(image_path).convert("RGB")
    size = list(image.size)
    declared = data.get("image_size")
    if declared is not None and list(declared) != size:
        die(f"image_size {declared} 与源图 {size} 不一致")
    fonts = installed_fonts()
    masks, texts, vectors = validate(data, size, fonts)

    if size[0] * PX < 914400 or size[1] * PX < 914400:
        die("PowerPoint 幻灯片每边至少 1 英寸，源图需至少 96×96 px")
    prs = Presentation()
    prs.slide_width = Emu(size[0] * PX)
    prs.slide_height = Emu(size[1] * PX)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(image_path), Emu(0), Emu(0), prs.slide_width, prs.slide_height)

    vector_dir = out_path.with_name(out_path.stem + "_vectors")
    results = []
    accepted = []
    for vec in vectors:
        x, y, w, h = [int(n) for n in vec["rect"]]
        crop = image.crop((x, y, x + w, y + h))
        kind = vec.get("kind", "geometry")
        try:
            bg, svg, err = trace_svg(crop)
        except Exception as exc:
            results.append({"id": vec["id"], "kind": kind, "accepted": False, "reason": str(exc)})
            continue
        if err:
            results.append({"id": vec["id"], "kind": kind, "accepted": False, "reason": err})
            continue
        rendered = composite(rasterize(svg, (w, h)), bg)
        ok, detail = matches(crop, rendered, bg)
        if not ok:
            results.append({"id": vec["id"], "kind": kind, "accepted": False, "reason": detail})
            continue
        vector_dir.mkdir(parents=True, exist_ok=True)
        svg_path = vector_dir / f"{vec['id']}.svg"
        svg_path.write_text(svg)
        fill = "#{:02X}{:02X}{:02X}".format(*bg)
        buf = io.BytesIO()
        rendered.save(buf, format="PNG")
        accepted.append((vec["id"], fill, [x, y, w, h], buf.getvalue(), svg.encode(), f"VECTOR {kind} {vec['id']}", str(svg_path), detail))
        results.append({"id": vec["id"], "kind": kind, "accepted": True, "svg": str(svg_path), "reason": detail})

    for mask in masks:
        add_mask(slide, mask["id"], [int(n) for n in mask["rect"]], mask["fill"])
    for item_id, fill, rect, _png, _svg, _name, _path, _detail in accepted:
        add_mask(slide, item_id, rect, fill)
    for text in texts:
        item = dict(text)
        item["rect"] = [int(n) for n in text["rect"]]
        add_text(slide, item)
    for _item_id, _fill, rect, png_bytes, svg_bytes, name, _path, _detail in accepted:
        add_svg_picture(slide, png_bytes, svg_bytes, rect, name)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    report = {"pptx": str(out_path), "image_size": size, "vectors": results}
    out_path.with_suffix(".result.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    return report


def self_check():
    from PIL import Image, ImageDraw

    fonts = installed_fonts()
    if not fonts:
        die("自检需要至少一种已安装字体")
    font = fonts[0]
    assert "/" not in font and ":" not in font
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        image = Image.new("RGB", (320, 120), (255, 255, 255))
        draw = ImageDraw.Draw(image)
        draw.line((140, 60, 250, 60), fill=(20, 40, 80), width=4)
        draw.polygon([(250, 52), (290, 60), (250, 68)], fill=(20, 40, 80))
        image.save(root / "fig.png")
        manifest = {
            "source_image": "fig.png",
            "image_size": [320, 120],
            "masks": [{"id": "label_bg", "rect": [8, 20, 80, 36], "fill": "#FFFFFF"}],
            "texts": [{
                "id": "label",
                "rect": [8, 20, 80, 36],
                "text": "Label",
                "font": font,
                "font_size": 18,
                "color": "#142850",
            }],
            "vectors": [{"id": "shaft", "kind": "arrow", "rect": [120, 40, 190, 48]}],
        }
        (root / "m.json").write_text(json.dumps(manifest))
        report = build(root / "m.json", root / "out.pptx")
        assert report["vectors"][0]["accepted"], report["vectors"][0]
        assert "<path" in Path(report["vectors"][0]["svg"]).read_text()
        with zipfile.ZipFile(root / "out.pptx") as zf:
            blob = "\n".join(zf.read(name).decode("utf-8", "replace") for name in zf.namelist())
            slide_xml = zf.read("ppt/slides/slide1.xml").decode("utf-8")
        assert "image/svg+xml" in blob
        assert "svgBlip" in blob
        assert "Label" in blob
        assert "effectRef" not in slide_xml
        assert "outerShdw" not in slide_xml
        crop = image.crop((120, 40, 310, 88))
        pale = crop.point(lambda v: min(255, v + 90))
        assert not matches(crop, pale, (255, 255, 255))[0], "a lighter trace must be rejected"
        bad = json.loads(json.dumps(manifest))
        bad["texts"][0]["font"] = "Definitely Missing Font XYZ"
        (root / "bad.json").write_text(json.dumps(bad))
        try:
            build(root / "bad.json", root / "bad.pptx")
        except SystemExit:
            pass
        else:
            raise AssertionError("missing font should exit")
    print("self-check ok")


def main(argv):
    require_libs()
    parser = argparse.ArgumentParser(description="AI PNG 到可编辑 PPTX")
    parser.add_argument("manifest", nargs="?")
    parser.add_argument("--out")
    parser.add_argument("--self-check", action="store_true")
    args = parser.parse_args(argv)
    if args.self_check:
        self_check()
        return
    if not args.manifest or not args.out:
        die("用法: build_overlay.py manifest.json --out figure.pptx")
    print(json.dumps(build(Path(args.manifest), Path(args.out)), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main(sys.argv[1:])
